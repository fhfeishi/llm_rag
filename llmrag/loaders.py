from __future__ import annotations

import html.parser
from pathlib import Path

from .models import DocumentPart, FileKind, LoadedDocument, SourceFile


TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".yaml", ".yml"}
CODE_SUFFIXES = {".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".go", ".rs", ".cpp", ".c", ".h", ".cs"}
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".webp"}


class _HTMLTextExtractor(html.parser.HTMLParser):
    def __init__(self):
        super().__init__()
        self._chunks: list[str] = []

    def handle_data(self, data: str) -> None:
        text = data.strip()
        if text:
            self._chunks.append(text)

    def text(self) -> str:
        return "\n".join(self._chunks)


def detect_file_kind(path: str | Path) -> FileKind:
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        return FileKind.pdf
    if suffix == ".pptx":
        return FileKind.pptx
    if suffix == ".xlsx":
        return FileKind.xlsx
    if suffix in IMAGE_SUFFIXES:
        return FileKind.image
    if suffix in {".html", ".htm"}:
        return FileKind.html
    if suffix in CODE_SUFFIXES:
        return FileKind.code
    if suffix in TEXT_SUFFIXES:
        return FileKind.text
    return FileKind.unknown


def load_file(path: str | Path) -> LoadedDocument:
    path = Path(path).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(path)

    kind = detect_file_kind(path)
    source = SourceFile(path=path, kind=kind)

    if kind in {FileKind.text, FileKind.code}:
        return _load_text_like(source)
    if kind == FileKind.html:
        return _load_html(source)
    if kind == FileKind.pdf:
        return _load_pdf(source)
    if kind == FileKind.pptx:
        return _load_pptx(source)
    if kind == FileKind.xlsx:
        return _load_xlsx(source)
    if kind == FileKind.image:
        return _load_image(source)
    raise ValueError(f"Unsupported file type: {path.suffix}")


def _load_text_like(source: SourceFile) -> LoadedDocument:
    text = source.path.read_text(encoding="utf-8", errors="ignore")
    part = DocumentPart(
        id="part-0001",
        title=source.path.name,
        text=text,
        source_path=source.path,
        kind=source.kind,
        index=1,
    )
    return LoadedDocument(source=source, title=source.path.stem, parts=[part])


def _load_html(source: SourceFile) -> LoadedDocument:
    raw = source.path.read_text(encoding="utf-8", errors="ignore")
    parser = _HTMLTextExtractor()
    parser.feed(raw)
    part = DocumentPart(
        id="part-0001",
        title=source.path.name,
        text=parser.text(),
        source_path=source.path,
        kind=FileKind.html,
        index=1,
    )
    return LoadedDocument(source=source, title=source.path.stem, parts=[part])


def _load_pdf(source: SourceFile) -> LoadedDocument:
    try:
        import PyPDF2
    except ImportError as exc:
        raise RuntimeError("PDF loading requires dependency: PyPDF2") from exc

    parts: list[DocumentPart] = []
    with source.path.open("rb") as f:
        reader = PyPDF2.PdfReader(f)
        for i, page in enumerate(reader.pages, 1):
            parts.append(
                DocumentPart(
                    id=f"page-{i:04d}",
                    title=f"Page {i}",
                    text=page.extract_text() or "",
                    source_path=source.path,
                    kind=FileKind.pdf,
                    index=i,
                    metadata={"page": i},
                )
            )
    return LoadedDocument(
        source=source,
        title=source.path.stem,
        parts=parts,
        metadata={"page_count": len(parts)},
    )


def _load_pptx(source: SourceFile) -> LoadedDocument:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("PPTX loading requires optional dependency: python-pptx") from exc

    prs = Presentation(str(source.path))
    parts: list[DocumentPart] = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                texts.append(shape.text)
        parts.append(
            DocumentPart(
                id=f"slide-{i:04d}",
                title=f"Slide {i}",
                text="\n".join(texts),
                source_path=source.path,
                kind=FileKind.pptx,
                index=i,
                metadata={"slide": i},
            )
        )
    return LoadedDocument(source=source, title=source.path.stem, parts=parts)


def _load_xlsx(source: SourceFile) -> LoadedDocument:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("XLSX loading requires optional dependency: openpyxl") from exc

    wb = openpyxl.load_workbook(source.path, data_only=True, read_only=True)
    parts: list[DocumentPart] = []
    for i, sheet in enumerate(wb.worksheets, 1):
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                rows.append("\t".join(values))
        parts.append(
            DocumentPart(
                id=f"sheet-{i:04d}",
                title=sheet.title,
                text="\n".join(rows),
                source_path=source.path,
                kind=FileKind.xlsx,
                index=i,
                metadata={"sheet": sheet.title},
            )
        )
    return LoadedDocument(source=source, title=source.path.stem, parts=parts)


def _load_image(source: SourceFile) -> LoadedDocument:
    part = DocumentPart(
        id="image-0001",
        title=source.path.name,
        source_path=source.path,
        kind=FileKind.image,
        index=1,
        metadata={"note": "Image bytes are not interpreted by the text loader yet."},
    )
    return LoadedDocument(source=source, title=source.path.stem, parts=[part])
